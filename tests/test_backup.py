import os
import shutil
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from src.backup import backup_to_db, export_to_xlsx, generate_full_backup, restore_from_db  # noqa: E402
from src.database import Database  # noqa: E402
from src.duplicate_detector import text_hash  # noqa: E402
from src.pipeline import JobHunterPipeline  # noqa: E402


class TestBackupRestore(unittest.TestCase):
    def setUp(self):
        self.tmp_dir = Path(tempfile.mkdtemp())
        self.db_path = str(self.tmp_dir / "job_hunter.db")
        self.backup_dir = self.tmp_dir / "backups"

        self.db = Database(self.db_path)
        self.db.upsert_user("user1", email="user1@example.com")
        self.db.set_setting("some_setting", "some_value")

        pipeline = JobHunterPipeline()
        record = pipeline.process(
            "Medical Representative — Al Rusafa, Baghdad. test@example.com",
            {"platform": "manual", "channel": None, "message_id": "1", "url": None, "date": "2026-08-30"},
        )
        self.db.insert_job(record, text_hash(record.raw_text))

    def tearDown(self):
        shutil.rmtree(self.tmp_dir, ignore_errors=True)

    def test_backup_to_db_creates_file(self):
        backup_path = backup_to_db(self.db_path, backup_dir=self.backup_dir)
        self.assertTrue(backup_path.exists())
        self.assertGreater(backup_path.stat().st_size, 0)

    def test_backup_contains_all_data(self):
        backup_path = backup_to_db(self.db_path, backup_dir=self.backup_dir)
        restored_db = Database(str(backup_path))
        user = restored_db.get_user("user1")
        self.assertEqual(user["email"], "user1@example.com")
        self.assertEqual(restored_db.get_setting("some_setting"), "some_value")
        jobs = restored_db.list_new_jobs(limit=10)
        self.assertEqual(len(jobs), 1)

    def test_restore_replaces_target_db_contents(self):
        backup_path = backup_to_db(self.db_path, backup_dir=self.backup_dir)

        # Simulate a fresh/different target DB (e.g. a new device):
        new_target_path = str(self.tmp_dir / "new_device.db")
        Database(new_target_path)  # creates empty schema

        restore_from_db(str(backup_path), new_target_path, keep_previous=False)

        restored = Database(new_target_path)
        user = restored.get_user("user1")
        self.assertIsNotNone(user)
        self.assertEqual(user["email"], "user1@example.com")
        self.assertEqual(len(restored.list_new_jobs(limit=10)), 1)

    def test_restore_keeps_previous_safety_copy_by_default(self):
        backup_path = backup_to_db(self.db_path, backup_dir=self.backup_dir)
        previous_copy = restore_from_db(str(backup_path), self.db_path, keep_previous=True)
        self.assertIsNotNone(previous_copy)
        self.assertTrue(previous_copy.exists())

    def test_export_to_xlsx_creates_readable_workbook(self):
        xlsx_path = export_to_xlsx(self.db_path, self.tmp_dir / "export.xlsx")
        self.assertTrue(xlsx_path.exists())

        from openpyxl import load_workbook
        wb = load_workbook(xlsx_path)
        self.assertIn("jobs", wb.sheetnames)
        self.assertIn("users", wb.sheetnames)
        jobs_sheet = wb["jobs"]
        self.assertGreaterEqual(jobs_sheet.max_row, 2)  # header + 1 data row

    def test_export_includes_every_table_in_the_live_schema(self):
        """
        H09 regression test: the export previously used a hardcoded table
        list that had drifted out of sync with the schema (missing
        job_applications, application_status_events, user_ignored_companies,
        user_saved_jobs, user_ignored_jobs). It now discovers tables
        directly from sqlite_master, so this can't silently drift again.
        """
        import sqlite3
        from openpyxl import load_workbook

        conn = sqlite3.connect(self.db_path)
        actual_tables = {
            r[0] for r in conn.execute(
                "SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'"
            ).fetchall()
        }
        conn.close()

        xlsx_path = export_to_xlsx(self.db_path, self.tmp_dir / "full_export.xlsx")
        wb = load_workbook(xlsx_path)
        exported_tables = set(wb.sheetnames)

        missing = actual_tables - exported_tables
        self.assertEqual(missing, set(), f"Tables missing from backup: {missing}")

    def test_generate_full_backup_produces_both_formats(self):
        paths = generate_full_backup(self.db_path, backup_dir=self.backup_dir)
        self.assertTrue(paths["db"].exists())
        self.assertTrue(paths["xlsx"].exists())


if __name__ == "__main__":
    unittest.main()
